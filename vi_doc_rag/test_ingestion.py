import os
import sys
import io
from contextlib import redirect_stdout
import tempfile
import unicodedata
import unittest
import importlib.util

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from vi_doc_rag import ChunkConfig, IngestionPipeline, clean_text, normalize_text
from vi_doc_rag.chunker import VietnameseChunker, Block, split_sentences, build_blocks
from vi_doc_rag.cleaner import remove_headers_footers, reflow_paragraphs

FONT = next((p for p in [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "arial.ttf"),
] if os.path.exists(p)), None)

REPORTLAB_AVAILABLE = importlib.util.find_spec("reportlab") is not None

PARA = ("Trí tuệ nhân tạo là lĩnh vực nghiên cứu các hệ thống có khả năng thực hiện tác vụ "
        "đòi hỏi trí thông minh của con người. Học máy là một nhánh quan trọng của lĩnh vực này. ")


BODIES = [
    "Phần đầu tiên nói về lịch sử hình thành của hệ thống thông tin.\nCác mốc quan trọng được trình bày theo thứ tự thời gian.\nMỗi mốc đi kèm một ví dụ minh hoạ.",
    "Phần tiếp theo mô tả kiến trúc tổng thể gồm nhiều thành phần.\nMỗi thành phần đảm nhiệm một chức năng riêng biệt.\nChúng giao tiếp với nhau qua giao diện chuẩn.",
    "Dữ liệu đầu vào được thu thập từ nhiều nguồn khác nhau.\nSau đó chúng được làm sạch và chuẩn hoá định dạng.\nBước này quyết định chất lượng của toàn bộ hệ thống.",
    "Việc đánh giá được thực hiện trên bộ dữ liệu kiểm thử độc lập.\nKết quả cho thấy độ chính xác cải thiện đáng kể.\nTuy vậy vẫn còn một số hạn chế cần khắc phục.",
    "Cuối cùng, báo cáo đề xuất hướng phát triển trong tương lai.\nTrọng tâm là mở rộng sang nhiều lĩnh vực mới.\nĐồng thời tối ưu chi phí vận hành hệ thống.",
]


class TestCleaner(unittest.TestCase):
    def test_nfc(self):
        decomposed = unicodedata.normalize("NFD", "Tiếng Việt")
        self.assertNotEqual(decomposed, "Tiếng Việt")
        self.assertEqual(normalize_text(decomposed), "Tiếng Việt")

    def test_junk_chars(self):
        self.assertEqual(normalize_text("Hà\u200b Nội\u00a0đẹp\x00"), "Hà Nội đẹp")

    def test_header_footer_removed(self):
        pages = [f"CÔNG TY ABC - BÁO CÁO\n{BODIES[i-1]}\nTrang {i}/5" for i in range(1, 6)]
        out = remove_headers_footers(pages)
        for p in out:
            self.assertNotIn("CÔNG TY ABC", p)
            self.assertNotIn("Trang", p)
            self.assertTrue(any(b in p for b in BODIES))

    def test_reflow(self):
        text = "Hệ thống này được xây dựng để phục vụ việc hỏi đáp trên\ntài liệu tiếng Việt một cách chính xác.\n- Mục một\n- Mục hai"
        out = reflow_paragraphs(text)
        self.assertIn("trên tài liệu", out)
        self.assertIn("\n- Mục một\n- Mục hai", out)

    def test_dot_leader(self):
        self.assertNotIn("....", clean_text("Chương 1 ............ 5"))


class TestChunker(unittest.TestCase):
    def test_sentence_abbrev(self):
        s = split_sentences("TS. Nguyễn Văn A sống ở TP. Hồ Chí Minh. Ông là giảng viên.")
        self.assertEqual(len(s), 2)

    def test_size_and_overlap(self):
        cfg = ChunkConfig(chunk_size=300, chunk_overlap=80)
        chunks = VietnameseChunker(cfg).split([Block(PARA * 8)], source="x.txt")
        self.assertGreater(len(chunks), 2)
        for c in chunks:
            self.assertLessEqual(len(c.text), cfg.chunk_size)
        for a, b in zip(chunks, chunks[1:]):  # có chồng lấp
            last_sentence = split_sentences(a.text)[-1]
            self.assertIn(last_sentence, b.text)

    def test_no_overlap_and_no_infinite_loop(self):
        chunks = VietnameseChunker(ChunkConfig(200, 0)).split([Block(PARA * 5)])
        self.assertTrue(all(len(c.text) <= 200 for c in chunks))

    def test_very_long_sentence(self):
        long = " ".join(["từ"] * 500)
        chunks = VietnameseChunker(ChunkConfig(200, 20)).split([Block(long)])
        self.assertTrue(all(len(c.text) <= 200 for c in chunks))
        self.assertEqual(sum(len(c.text.split()) for c in chunks) >= 500, True)

    def test_heading_section_metadata(self):
        blocks = [Block("Chương 1 Giới thiệu", is_heading=True), Block(PARA * 3),
                  Block("Chương 2 Phương pháp", is_heading=True), Block(PARA * 3)]
        chunks = VietnameseChunker(ChunkConfig(500, 50, min_chunk_size=100)).split(blocks)
        secs = {c.metadata["section"] for c in chunks}
        self.assertEqual(secs, {"Chương 1 Giới thiệu", "Chương 2 Phương pháp"})

    def test_heading_detected_inside_multiline_paragraph(self):
        blocks = build_blocks(["Giới thiệu\nChương 1. Mục tiêu\nNội dung chương."], has_pages=False)
        self.assertTrue(any(b.is_heading and b.text == "Chương 1. Mục tiêu" for b in blocks))
    def test_cross_page_merge(self):
        pages = ["Đây là câu bị ngắt giữa hai trang và chưa kết thúc", "tiếp tục ở trang sau. Câu tiếp theo."]
        blocks = build_blocks(pages, has_pages=True)
        self.assertEqual(len(blocks), 1)
        self.assertEqual((blocks[0].page_start, blocks[0].page_end), (1, 2))


class TestPipeline(unittest.TestCase):
    def test_docx(self):
        import docx
        d = docx.Document()
        d.add_heading("Giới thiệu hệ thống", level=1)
        d.add_paragraph(PARA * 3)
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Tên", "Giá trị"
        t.cell(1, 0).text, t.cell(1, 1).text = "Phiên bản", "1.0"
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "a.docx")
            d.save(p)
            chunks = IngestionPipeline().process_file(p)
        full = "\n".join(c.text for c in chunks)
        self.assertIn("Trí tuệ nhân tạo", full)
        self.assertIn("Tên | Giá trị", full)
        self.assertEqual(chunks[0].metadata["section"], "Giới thiệu hệ thống")

    @unittest.skipUnless(FONT and REPORTLAB_AVAILABLE, "Cần font tiếng Việt và reportlab để chạy test PDF")
    def test_pdf_header_footer_pages(self):
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
        pdfmetrics.registerFont(TTFont("VN", FONT))
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "b.pdf")
            c = canvas.Canvas(p)
            for i in range(1, 5):
                c.setFont("VN", 9)
                c.drawString(60, 800, "TRƯỜNG ĐẠI HỌC XYZ - GIÁO TRÌNH")
                c.drawString(280, 30, f"Trang {i}")
                c.setFont("VN", 11)
                y = 740
                for line in BODIES[i - 1].split("\n"):
                    if line.strip():
                        c.drawString(60, y, line.strip()); y -= 16
                c.showPage()
            c.save()
            chunks = IngestionPipeline(ChunkConfig(400, 50)).process_file(p)
        full = "\n".join(c.text for c in chunks)
        self.assertNotIn("GIÁO TRÌNH", full)
        self.assertNotIn("Trang 2", full)
        self.assertIn("lịch sử hình thành", full)
        self.assertIn("giao diện chuẩn", full)
        self.assertEqual(max(c.metadata["page_end"] for c in chunks), 4)

    def test_txt_and_jsonl(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = os.path.join(tmp, "c.txt")
            with open(p, "w", encoding="utf-8") as f:
                f.write(PARA * 6)
            pipe = IngestionPipeline(ChunkConfig(300, 50))
            chunks = pipe.process_directory(tmp)
            out = os.path.join(tmp, "o.jsonl")
            pipe.save_jsonl(chunks, out)
            with open(out, encoding="utf-8") as fh:
                self.assertEqual(len(fh.readlines()), len(chunks))
        self.assertEqual(len({c.id for c in chunks}), len(chunks))


class TestCLI(unittest.TestCase):
    def test_cli_runs_on_path(self):
        import subprocess
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        with tempfile.TemporaryDirectory() as tmp:
            src, out = os.path.join(tmp, "d.txt"), os.path.join(tmp, "o.jsonl")
            with open(src, "w", encoding="utf-8") as f:
                f.write(PARA * 4)
            r = subprocess.run([sys.executable, "-m", "vi_doc_rag", src, "--out", out],
                               cwd=root, capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertIn("Tổng", r.stdout)
            self.assertTrue(os.path.exists(out))

    def test_picker_falls_back_without_display(self):
        from unittest import mock
        from vi_doc_rag import picker
        with mock.patch.dict(sys.modules, {"tkinter": None}), \
                mock.patch("builtins.input", return_value='  "C:\\a b\\x.pdf" '):
            with redirect_stdout(io.StringIO()):
                selected = picker.pick_path()
            self.assertEqual(selected, "C:\\a b\\x.pdf")


if __name__ == "__main__":
    unittest.main(verbosity=2)





