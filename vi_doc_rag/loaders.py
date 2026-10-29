"""Đọc file PDF / DOCX / TXT / MD thành RawDocument."""
from __future__ import annotations

import logging
import os
from typing import List

from .models import RawDocument

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
_MIN_CHARS_PER_PAGE = 20  # dưới ngưỡng này coi như trang không có text


class UnsupportedFileError(ValueError):
    pass


def load_document(path: str) -> RawDocument:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in (".txt", ".md"):
        return load_text(path)
    raise UnsupportedFileError(
        f"Định dạng '{ext}' chưa được hỗ trợ. Hỗ trợ: {sorted(SUPPORTED_EXTENSIONS)}"
    )


# --------------------------------------------------------------------- PDF
def load_pdf(path: str) -> RawDocument:
    """Đọc PDF bằng pypdf; trang nào trống thì thử lại bằng pdfplumber.

    Lưu ý: PDF dạng ảnh scan không có lớp text -> cần OCR (chưa làm ở bước này).
    """
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError(f"PDF bị mã hoá, không mở được: {path}") from exc

    pages: List[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            logger.warning("pypdf lỗi ở 1 trang của %s: %s", path, exc)
            pages.append("")

    # Fallback pdfplumber cho các trang trống
    weak = [i for i, t in enumerate(pages) if len(t.strip()) < _MIN_CHARS_PER_PAGE]
    if weak:
        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                for i in weak:
                    if i < len(pdf.pages):
                        alt = pdf.pages[i].extract_text(x_tolerance=1.5) or ""
                        if len(alt.strip()) > len(pages[i].strip()):
                            pages[i] = alt
        except Exception as exc:  # noqa: BLE001
            logger.warning("Fallback pdfplumber thất bại: %s", exc)

    empty_pages = [i + 1 for i, t in enumerate(pages) if len(t.strip()) < _MIN_CHARS_PER_PAGE]
    if empty_pages and len(empty_pages) == len(pages):
        logger.warning("%s: không trích xuất được text (có thể là PDF scan, cần OCR).", path)
    elif empty_pages:
        logger.info("%s: các trang gần như trống: %s", path, empty_pages)

    return RawDocument(
        source=path,
        pages=pages,
        has_pages=True,
        metadata={"format": "pdf", "num_pages": len(pages), "empty_pages": empty_pages},
    )


# -------------------------------------------------------------------- DOCX
def load_docx(path: str) -> RawDocument:
    """Đọc DOCX theo đúng thứ tự xuất hiện của đoạn văn và bảng.

    - Heading -> "# Tiêu đề" (để chunker nhận ra mục).
    - Bảng -> mỗi hàng 1 dòng "ô1 | ô2 | ô3".
    - Các đoạn cách nhau bằng dòng trống.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    blocks: List[str] = []

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            low = style.lower()
            if low.startswith(("heading", "title", "tiêu đề")):
                text = "# " + text
            elif "list" in low and not text.startswith(("-", "•", "+", "*")):
                text = "- " + text
            blocks.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # ô bị gộp xuất hiện lặp lại
                        continue
                    seen.add(id(cell._tc))
                    t = " ".join(cell.text.split())
                    if t:
                        cells.append(t)
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append("\n".join(rows))

    return RawDocument(
        source=path,
        pages=["\n\n".join(blocks)],
        has_pages=False,
        metadata={"format": "docx"},
    )


# --------------------------------------------------------------- TXT / MD
def load_text(path: str) -> RawDocument:
    for enc in ("utf-8-sig", "utf-16", "cp1258"):  # cp1258 = bảng mã Windows tiếng Việt
        try:
            with open(path, "r", encoding=enc) as f:
                text = f.read()
            return RawDocument(source=path, pages=[text], has_pages=False,
                               metadata={"format": os.path.splitext(path)[1][1:], "encoding": enc})
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise RuntimeError(f"Không đoán được encoding của file: {path}")
